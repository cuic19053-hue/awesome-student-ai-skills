#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""docx_layout_fix.py — 对已生成的 .docx 做一次排版硬化
====================================================

**存在的理由**：35 个子 skill 的 build.py 各自手写 python-docx 代码，
统计下来没有一个是「排版安全」的——

    cantSplit（表格不可拆行）      0 / 35
    keep_with_next（图注绑定）     0 / 35
    tblHeader（长表表头重复）      0 / 35
    outlineLvl（标题大纲级别）     0 / 35
    core_properties（文档属性）    0 / 35
    表题「表 N」                   1 / 35

内容写得再好，排版裸奔一样会被评审一眼看出是模板套的。与其改 35 份
build.py（每份写法都不同，改完还得逐份回归），不如在**产物层**统一硬化：
build.py 保存文件后调一次 `fix(path)`，或者用户拿任意 .docx 跑一遍 CLI。

能修什么（对应 references/layout_standard.md 的条款）
---------------------------------------------------
    L06 页码     无 PAGE 域则补「第 N 页　共 M 页」，封面不显示
    L07 表格     每行 cantSplit；≤13 行的表整表不跨页
    L08 表头     行数 ≥8 的表，首行 tblHeader 跨页重复
    L10 图注     图片段落 keep_with_next + keep_together
    L11 属性     清掉 author="python-docx" 这类模板痕迹
    L13 结构     章/节标题补 outlineLvl，导航窗格与自动目录可用

**修不了什么**（必须由内容生成侧负责）
------------------------------------
    L04 图表引用闭环 —— 需要判断「哪句话该提到表 3」，机器猜不准
    L09 表题        —— 需要知道这张表在讲什么
    L03 编号连续    —— 同上
    这两类靠 utils/layout_audit.py 报出来，由作者/agent 补。

用法
----
    from utils.docx_layout_fix import fix
    fix("out.docx")                 # 就地硬化
    fix("out.docx", "fixed.docx")   # 另存

    python utils/docx_layout_fix.py out.docx
    python utils/docx_layout_fix.py examples/demos --report

作者：awesome-student-ai-skills
版本：1.0.0
"""

from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path
from typing import List, Optional, Sequence

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

H1 = re.compile(r"^[一二三四五六七八九十]+、")
H2 = re.compile(r"^\d+\.\d+")
CAP = re.compile(r"^([图表])\s*\d+")

BIG_TABLE_ROWS = 8          # ≥ 此行数：表头跨页重复
KEEP_TOGETHER_MAX_ROWS = 13  # ≤ 此行数：整表不跨页


# ============================================================
# 单项修复
# ============================================================

def fix_tables(doc) -> int:
    """表格三道保护：行级不拆、长表表头重复、短表整表不跨页。返回处理的表格数。"""
    n = 0
    for table in doc.tables:
        rows = table.rows
        if len(rows) <= 1:
            continue
        n += 1
        for row in rows:
            trPr = row._tr.get_or_add_trPr()
            if trPr.find(qn("w:cantSplit")) is None:
                trPr.append(OxmlElement("w:cantSplit"))
        if len(rows) >= BIG_TABLE_ROWS:
            trPr = rows[0]._tr.get_or_add_trPr()
            if trPr.find(qn("w:tblHeader")) is None:
                trPr.append(OxmlElement("w:tblHeader"))
        if len(rows) <= KEEP_TOGETHER_MAX_ROWS:
            for row in rows[:-1]:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        p.paragraph_format.keep_with_next = True
                        p.paragraph_format.keep_together = True
    return n


def fix_figures(doc) -> int:
    """图片段与下一段（图注）绑定，避免图在页底、图注跑下页。"""
    n = 0
    for p in doc.paragraphs:
        if not p._p.findall(".//" + qn("w:drawing")):
            continue
        n += 1
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.keep_together = True
    return n


def fix_headings(doc) -> int:
    """给章/节标题补大纲级别。只认「加粗 + 序号开头」的段落，避免误伤正文。"""
    n = 0
    for p in doc.paragraphs:
        t = p.text.strip()
        if not t or not p.runs or not p.runs[0].bold:
            continue
        lvl = 0 if H1.match(t) else (1 if H2.match(t) else None)
        if lvl is None:
            continue
        pPr = p._p.get_or_add_pPr()
        if pPr.find(qn("w:outlineLvl")) is not None:
            continue
        el = OxmlElement("w:outlineLvl")
        el.set(qn("w:val"), str(lvl))
        pPr.insert_element_before(el, "w:rPr", "w:sectPr")
        n += 1
    return n


def _has_page_field(doc) -> bool:
    for sec in doc.sections:
        for p in sec.footer.paragraphs:
            if "PAGE" in p._p.xml:
                return True
    return False


def _field(p, instr: str, placeholder: str = "1", size: int = 9):
    def _fc(kind):
        r = p.add_run()
        el = OxmlElement("w:fldChar")
        el.set(qn("w:fldCharType"), kind)
        r._element.append(el)
        r.font.size = Pt(size)
        return r

    _fc("begin")
    r = p.add_run()
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = f" {instr} "
    r._element.append(it)
    r.font.size = Pt(size)
    _fc("separate")
    rv = p.add_run(placeholder)
    rv.font.size = Pt(size)
    _fc("end")


def fix_page_number(doc) -> bool:
    """没有页码就补「第 N 页　共 M 页」，并让封面不显示。返回是否改动。"""
    if _has_page_field(doc):
        return False
    for sec in doc.sections:
        sec.different_first_page_header_footer = True
        footer = sec.footer
        p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        for r in list(p.runs):
            r.text = ""
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run("第 ").font.size = Pt(9)
        _field(p, "PAGE")
        p.add_run(" 页　共 ").font.size = Pt(9)
        _field(p, "NUMPAGES")
        p.add_run(" 页").font.size = Pt(9)
    return True


def fix_properties(doc) -> bool:
    cp = doc.core_properties
    author = (cp.author or "").strip()
    if author and "python-docx" not in author.lower():
        return False
    cp.author = "文档助手"
    cp.last_modified_by = "文档助手"
    return True


# ============================================================
# 组合
# ============================================================

def fix(path, out=None, verbose: bool = False) -> dict:
    """对 .docx 做一次排版硬化，返回各项改动量。"""
    src = Path(path)
    doc = Document(str(src))

    report = {
        "tables": fix_tables(doc),
        "figures": fix_figures(doc),
        "headings": fix_headings(doc),
        "page_number": fix_page_number(doc),
        "properties": fix_properties(doc),
    }

    dst = Path(out) if out else src
    doc.save(str(dst))
    report["saved"] = str(dst)

    if verbose:
        print(f"[layout_fix] {src.name}")
        print(f"  表格保护 {report['tables']} 张 · 图注绑定 {report['figures']} 处 · "
              f"标题大纲 {report['headings']} 个")
        print(f"  页码 {'已补' if report['page_number'] else '已有'} · "
              f"文档属性 {'已修' if report['properties'] else '正常'}")
        print(f"  → {dst}")
    return report


def main(argv: Optional[Sequence[str]] = None) -> int:
    ap = argparse.ArgumentParser(description="对 docx 做排版硬化（表格防拆 / 图注绑定 / 页码 / 大纲）")
    ap.add_argument("targets", nargs="+", help="docx 文件或目录")
    ap.add_argument("--out", help="另存路径（仅单文件时可用）")
    ap.add_argument("--report", action="store_true", help="打印改动明细")
    args = ap.parse_args(argv)

    files: List[Path] = []
    for t in args.targets:
        p = Path(t)
        if p.is_dir():
            files.extend(sorted(p.rglob("*.docx")))
        elif p.suffix.lower() == ".docx":
            files.append(p)
    files = [f for f in files if not f.name.startswith("~$")]
    if not files:
        print("未找到 .docx", file=sys.stderr)
        return 2
    if args.out and len(files) > 1:
        print("--out 只能用于单个文件", file=sys.stderr)
        return 2

    total = {"tables": 0, "figures": 0, "headings": 0}
    for f in files:
        rep = fix(f, out=args.out, verbose=args.report)
        for k in total:
            total[k] += rep[k]
    print(f"完成 {len(files)} 份 · 表格保护 {total['tables']} 张 · "
          f"图注绑定 {total['figures']} 处 · 标题大纲 {total['headings']} 个")
    return 0


if __name__ == "__main__":
    sys.exit(main())
