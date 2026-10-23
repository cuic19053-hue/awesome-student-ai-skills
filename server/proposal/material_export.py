"""Production DOCX renderer for the innovation pilot (no demonstration data)."""
from __future__ import annotations

import re
import shutil
import subprocess
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from .materials import material_view

STYLE = {
    "paper": "A4", "top_cm": 2.5, "bottom_cm": 2.5, "left_cm": 3.0, "right_cm": 2.6,
    "body_pt": 12, "line_spacing": 1.5, "first_line_pt": 24,
    "font": "Noto Serif CJK SC", "heading_font": "Noto Sans CJK SC",
    "figure_width_cm": 15.0,
    "cover_org_pt": 22, "cover_kind_pt": 26, "cover_title_pt": 16,
    "heading1_pt": 15, "heading2_pt": 12.5, "caption_pt": 10.5,
}


def blocks(content: str) -> list[dict[str, Any]]:
    """A lossless-enough typed view; raw section text remains authoritative."""
    lines, result, i = content.splitlines(), [], 0
    while i < len(lines):
        line = lines[i].strip()
        if line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [x.strip() for x in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r"[:\- ]*", x) for x in cells):
                    rows.append(cells)
                i += 1
            if rows:
                result.append({"type": "table", "rows": rows})
            continue
        if line:
            match = re.fullmatch(r"\[\[figure:([a-zA-Z0-9_-]+)\]\]", line)
            if match:
                result.append({"type": "figure", "id": match[1]})
            else:
                result.append({"type": "heading" if line.startswith("#") else "paragraph",
                               "text": re.sub(r"^#{1,6}\s*", "", line)})
        i += 1
    return result


def _font(run, size=12, bold=False, family=None):
    family = family or (STYLE["heading_font"] if bold else STYLE["font"])
    run.font.name = family
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), family)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)


def _text(container, text, *, heading=False, size=None):
    p = container.add_paragraph()
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(6)
    if heading:
        p.paragraph_format.keep_with_next = True
    else:
        p.paragraph_format.first_line_indent = Pt(24)
    for index, segment in enumerate(re.split(r"\*\*(.*?)\*\*", text)):
        body_family = container.part.document.styles["Normal"].font.name or STYLE["font"]
        _font(p.add_run(segment), size=size or (15 if heading else 12), bold=heading or index % 2 == 1,
              family=STYLE["heading_font"] if heading else body_family)
    return p


def chapter_number(value: int) -> str:
    digits = "零一二三四五六七八九"
    if value < 10:
        return digits[value]
    return (digits[value // 10] if value >= 20 else "") + "十" + (digits[value % 10] if value % 10 else "")


class ParagraphRegion:
    """Insert blocks at a template anchor, rather than appending a duplicate chapter."""
    def __init__(self, parent, paragraph):
        self.parent = parent
        self.cursor = paragraph._p
        self.part = parent.part

    def add_paragraph(self):
        paragraph = self.parent.add_paragraph()
        self.cursor.addnext(paragraph._p)
        self.cursor = paragraph._p
        return paragraph

    def add_table(self, rows, cols):
        table = self.parent.add_table(rows=rows, cols=cols)
        self.cursor.addnext(table._tbl)
        self.cursor = table._tbl
        return table


class BodyRegion:
    """Append generated content immediately before template signature/review blocks."""
    def __init__(self, document, anchor):
        self.parent = document
        self.anchor = anchor
        self.part = document.part

    def _insert(self, element):
        element.getparent().remove(element)
        self.anchor.addprevious(element)

    def add_paragraph(self):
        p = self.parent.add_paragraph()
        self._insert(p._p)
        return p

    def add_table(self, rows, cols):
        table = self.parent.add_table(rows=rows, cols=cols)
        self._insert(table._tbl)
        return table


def export_material(store, project_id: str, folder: Path) -> dict[str, Any]:
    from .exporter import _fill_labeled_tables, _fill_labeled_paragraphs, _replace_tokens, _match_section_id
    folder.mkdir(parents=True, exist_ok=True)
    project = store.project_for_export(project_id) if hasattr(store, "project_for_export") else store.get_project(project_id)
    docs = store.list_documents(project_id)
    project["_documents"] = docs
    view = material_view(project, docs)
    effective = view["requirements"]["effective"]
    template = next((d for d in docs if d["id"] == project.get("template_document_id")), None)
    warnings, containers, token_sections, paragraph_anchors = [], {}, set(), {}
    used_template = bool(template and Path(template["stored_path"]).suffix.lower() == ".docx")
    if used_template:
        document = Document(template["stored_path"])
        regions = [(document, p) for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    regions.extend((cell, p) for p in cell.paragraphs)
        for parent, paragraph in regions:
            for sid in re.findall(r"\{\{section:([a-zA-Z0-9_-]+)\}\}", paragraph.text):
                if sid not in project["sections"]:
                    continue
                token_sections.add(sid)
                paragraph_anchors[sid] = (parent, paragraph)
        # Locate original template regions before replacing their labels.
        for table in document.tables:
            for row in table.rows:
                if len(row.cells) >= 2:
                    sid = _match_section_id(row.cells[0].text, project)
                    if sid:
                        containers[sid] = row.cells[1]
        replaced = _replace_tokens(document, project)
        filled = _fill_labeled_tables(document, project) | _fill_labeled_paragraphs(document, project) | token_sections
        for sid in token_sections:
            parent, paragraph = paragraph_anchors[sid]
            # Keep template paragraph properties; replace only generated body.
            residual = paragraph.text.replace(str(project["sections"].get(sid) or "待补充"), "", 1)
            paragraph.clear()
            if residual:
                paragraph.add_run(residual)
            containers[sid] = ParagraphRegion(parent, paragraph)
        if not replaced and not filled:
            warnings.append("学校模板未找到可映射区域，原件保留，正文附在后面；模板适配未完成")
        for sid in filled:
            if sid not in containers:
                warnings.append(f"章节 {sid} 已填入段落区域，附图单独保留，请核对模板位置")
    else:
        document, filled = Document(), set()
        section = document.sections[0]
        section.page_width, section.page_height = Cm(21), Cm(29.7)
        section.top_margin, section.bottom_margin = Cm(2.5), Cm(2.5)
        section.left_margin, section.right_margin = Cm(3), Cm(2.6)
        normal = document.styles["Normal"]
        normal.font.name, normal.font.size = STYLE["font"], Pt(12)
        normal._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), STYLE["font"])
        font_rule = effective.get("font")
        if font_rule:
            requested = next((name for name in ("宋体", "黑体", "楷体", "仿宋", "Noto Serif CJK SC", "Noto Sans CJK SC")
                              if name in font_rule["value"]), None)
            if requested:
                normal.font.name = requested
                normal._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), requested)
                warnings.append(f"要求字体为 {requested}，需核对渲染环境是否发生字体替代")
            else:
                warnings.append("字体要求无法自动解释，已保留原文要求，请人工核对")
        cover_org = _text(document, "大学生创新创业训练计划", heading=True, size=22)
        cover_org.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cover_org.paragraph_format.space_before = Pt(28)
        cover_org.paragraph_format.space_after = Pt(20)
        cover_kind = _text(document, "创新训练项目申报书（草稿）", heading=True, size=26)
        cover_kind.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cover_kind.paragraph_format.space_after = Pt(24)
        _text(document, project["title"], heading=True, size=16).alignment = WD_ALIGN_PARAGRAPH.CENTER
        info = document.add_table(rows=0, cols=2)
        info.style = "Table Grid"
        for label, value in (
            ("项目类别", "创新训练项目"), ("申报级别", project.get("level")),
            ("所属学科", project.get("discipline")), ("申报学校", project.get("school")),
            ("项目负责人", project["state"].get("leader_name")),
            ("指导教师", project["state"].get("advisor_name")),
            ("项目周期", project["state"].get("duration")),
            ("申请经费", project["state"].get("budget")),
        ):
            cells = info.add_row().cells
            cells[0].text, cells[1].text = label, value or "待补充"
        info.alignment = 1
        for row in info.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_before = Pt(3)
                    paragraph.paragraph_format.space_after = Pt(3)
        _text(document, "草稿：请核对正式通知、事实、引用及待补充项。")
        document.add_page_break()
        if template:
            warnings.append("非 Word 模板不能保留原版式；使用可编辑通用排版，需人工核对")
    # Keep approval/signature blocks last. Adding unmapped chapters after the
    # signature area falsely presents them as part of the official form.
    body_container = document
    if used_template:
        approval = ("审核意见", "评审意见", "审批意见", "学院意见", "学校意见",
                    "申请人签字", "负责人签字", "指导教师签字", "单位盖章")
        anchor = next((p._p for p in document.paragraphs
                       if any(marker in p.text for marker in approval)), None)
        if anchor is None:
            anchor = next((table._tbl for table in document.tables
                           if any(marker in cell.text for row in table.rows for cell in row.cells
                                  for marker in approval)), None)
        if anchor is not None:
            body_container = BodyRegion(document, anchor)

    if used_template:
        physical_order = {node: index for index, node in enumerate(document._element.body.iter())}
        def template_rank(section):
            region = containers.get(section["id"])
            element = region.cursor if isinstance(region, ParagraphRegion) else getattr(region, "_tc", None)
            return physical_order.get(element, len(physical_order) + 1)
        project["outline"] = sorted(project["outline"], key=template_rank)
    def figure_position(figure):
        sid = figure.get("section_id")
        content_blocks = blocks(project["sections"].get(sid, ""))
        explicit = next((i for i, block in enumerate(content_blocks)
                         if block["type"] == "figure" and block["id"] == figure["id"]), None)
        anchor = figure.get("anchor", -1)
        position = explicit if explicit is not None else (
            anchor + .5 if 0 <= anchor < len(content_blocks) else len(content_blocks) + 1)
        return (next((i for i, s in enumerate(project["outline"]) if s["id"] == sid), 1000),
                position, figure["id"])
    figures = sorted(project.get("diagrams", []), key=figure_position)
    numbers = {d["id"]: n for n, d in enumerate(figures, 1)}
    inserted = set()
    table_number = 0

    def references(text):
        def replace(match):
            key = match[1]
            if key not in numbers:
                warnings.append(f"图表引用不存在：{key}")
                return f"【缺失图表：{key}】"
            return f"图 {numbers[key]}"
        return re.sub(r"\[\[ref:([a-zA-Z0-9_-]+)\]\]", replace, text)

    def insert_figure(container, figure):
        key = figure["id"]
        if key in inserted:
            return
        inserted.add(key)
        number = numbers[key]
        path = Path(figure.get("preview") or "")
        _text(container, f"相关图示见图 {number}。")
        if path.is_file():
            from PIL import Image
            with Image.open(path) as picture:
                width_cm = 13 if container is not document else 15
                too_tall = width_cm * picture.height / picture.width > 19
            if too_tall:
                warnings.append(f"图 {number} 过高，请拆分后插入；未通过缩小到不可读字号解决")
                _text(container, f"【待补充：图 {number} 请拆分，原图保留在项目中】")
            else:
                p = container.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.keep_with_next = True
                p.add_run().add_picture(str(path), width=Cm(width_cm))
        else:
            warnings.append(f"图 {number} 图片缺失")
            _text(container, f"【待补充：图 {number} 图片缺失】")
        caption = _text(container, f"图 {number}　{figure.get('title', '项目图示')}", size=10.5)
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption.paragraph_format.first_line_indent = Pt(0)
        source = figure.get("source") or ("AI 生成概念示意，不代表实物或实测结果" if figure.get("renderer") == "gpt" else "项目资料，待核对")
        _text(container, f"来源：{source}")
        if figure.get("scene_text") and not figure.get("text_confirmed"):
            warnings.append(f"图 {number} 场景文字尚未人工核对")

    for index, section in enumerate(project["outline"], 1):
        sid = section["id"]
        associated = [d for d in figures if d.get("section_id") == sid]
        if sid in filled:
            container = containers.get(sid)
            if container is None:
                for figure in associated:
                    insert_figure(body_container, figure)
                continue
            if not isinstance(container, ParagraphRegion):
                container.text = ""  # Only cells actually filled from a blank/placeholder.
        else:
            container = body_container
            heading = _text(container, f"{chapter_number(index)}、{section['title']}", heading=True)
            heading.style = document.styles["Heading 1"]
        content = project["sections"].get(sid) or "【待补充：本章内容】"
        section_blocks = blocks(content)
        explicit_ids = {b["id"] for b in section_blocks if b["type"] == "figure"}
        for block_index, block in enumerate(section_blocks):
            if block["type"] == "figure":
                figure = next((d for d in associated if d["id"] == block["id"]), None)
                if figure:
                    insert_figure(container, figure)
                else:
                    warnings.append(f"本章图表占位不存在或不属于本章：{block['id']}")
            elif block["type"] == "table":
                table_number += 1
                _text(container, f"本章数据见表 {table_number}。")
                cap = _text(container, f"表 {table_number}　本章资料表", size=10.5)
                cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                cap.paragraph_format.first_line_indent = Pt(0)
                cap.paragraph_format.keep_with_next = True
                rows = block["rows"]
                table = container.add_table(rows=0, cols=max(map(len, rows)))
                table.style = "Table Grid"
                for row in rows:
                    cells = table.add_row().cells
                    for k, value in enumerate(row):
                        cells[k].text = references(value)
            else:
                _text(container, references(block["text"]), heading=block["type"] == "heading",
                      size=12.5 if block["type"] == "heading" else None)
            for figure in associated:
                if figure["id"] not in explicit_ids and figure.get("anchor", -1) == block_index:
                    insert_figure(container, figure)
        for figure in associated:
            insert_figure(container, figure)
    for figure in figures:
        if figure["id"] not in inserted:
            warnings.append(f"图表 {figure['id']} 尚未关联有效章节，保留在附录")
            insert_figure(document, figure)
    if any(re.search(r"待补充|待验证|\{\{", str(v)) for v in project["sections"].values()):
        warnings.append("正文仍包含待补充或待验证内容")
    if view["stale"]:
        warnings.append("存在可能过期的章节或图表")
    if any(x["status"] == "needs_confirmation" for x in view["requirements"]["conflicts"]):
        warnings.append("存在未解决的同级要求冲突")
    for conflict in view["requirements"]["conflicts"]:
        if conflict["status"] == "priority_applied":
            warnings.append(f"要求“{conflict['key']}”存在不同优先级内容；草稿按已记录优先级处理，请核对")
    if used_template and effective.get("font"):
        warnings.append(f"学校模板优先；正式字体要求为“{effective['font']['value']}”，未静默覆盖模板字体")
    target = folder / "proposal.docx"
    for key in ("word_count", "page_limit", "format", "sections", "content"):
        if key in effective and not (used_template and key == "font"):
            warnings.append(f"请核对适用要求 {key}：{effective[key]['value']}")
    document.core_properties.title = project["title"]
    document.core_properties.author = ""
    document.save(target)
    from utils.docx_layout_fix import fix
    # Preserve formal templates: the generic hardener can alter their pagination.
    if not used_template:
        fix(target)
    from .tools import ProposalTools
    validation = ProposalTools(store).validate_project(project_id)
    return {"file": "proposal.docx", "warnings": sorted(set(warnings)), "validation": validation,
            "effective_requirements": effective,
            "template_status": "mapped" if filled or (used_template and replaced) else ("needs_review" if template else "generic"),
            "input_snapshot": view["snapshot"], "layout_profile": STYLE}


def check_pages(folder: Path, report: dict[str, Any]) -> dict[str, Any]:
    """Actual conversion gate; failed/skipped checks are never reported passed."""
    from utils.layout_audit import audit
    from .material_figures import font_path
    static = audit(folder / "proposal.docx", render=False)
    checks = [{"rule": x.rule, "status": "failed" if x.level == "FAIL" else "needs_review",
               "detail": x.detail} for x in static.issues]
    checks.append({"rule": "font", "status": "needs_review",
                   "detail": "须核对实际字体替代与特殊字形；不能仅凭字体配置断言完整"})
    engine = shutil.which("soffice") or shutil.which("libreoffice")
    if not engine:
        checks.append({"rule": "page_render", "status": "not_checked", "detail": "未安装 LibreOffice"})
        return {**report, "checks": checks, "quality_status": "not_checked", "pages": []}
    font_path()  # Explicit failure on missing CJK font.
    profile = folder / "lo-profile"
    process = subprocess.run(
        [engine, f"-env:UserInstallation={profile.resolve().as_uri()}", "--headless", "--norestore",
         "--convert-to", "pdf", "--outdir", str(folder.resolve()), str((folder / "proposal.docx").resolve())],
        capture_output=True, timeout=180,
    )
    pdf = folder / "proposal.pdf"
    if process.returncode or not pdf.is_file():
        checks.append({"rule": "page_render", "status": "failed", "detail": "Word 转 PDF 失败"})
        return {**report, "checks": checks, "quality_status": "failed", "pages": []}
    import pypdfium2 as pdfium
    doc = pdfium.PdfDocument(str(pdf))
    if len(doc) > 160:
        doc.close()
        raise ValueError("文档超过 160 页，请拆分后检查")
    pages = []
    for index in range(len(doc)):
        page = doc[index]
        textpage = page.get_textpage()
        text = textpage.get_text_range()
        if "\ufffd" in text or "\u25a1" in text:
            checks.append({"rule": "glyph", "status": "needs_review", "detail": f"第 {index+1} 页可能缺字"})
        if not text.strip() and not list(page.get_objects()):
            checks.append({"rule": "blank_page", "status": "needs_review", "detail": f"第 {index+1} 页可能为空白"})
        name = f"page-{index+1:03d}.png"
        bitmap = page.render(scale=1.25)
        bitmap.to_pil().save(folder / name)
        bitmap.close()
        pages.append(name)
        textpage.close()
        page.close()
    doc.close()
    checks.append({"rule": "page_render", "status": "passed", "detail": f"已渲染 {len(pages)} 页"})
    try:
        from utils.layout_audit import check_rendered
        before_count = len(static.issues)
        check_rendered(static, pdf)
        for issue in static.issues[before_count:]:
            checks.append({"rule": issue.rule, "status": "needs_review", "detail": issue.detail})
        if any("pymupdf" in s for s in static.skipped):
            checks.append({"rule": "page_geometry", "status": "not_checked", "detail": "未安装 PyMuPDF"})
    except Exception:
        checks.append({"rule": "page_geometry", "status": "not_checked", "detail": "页面几何检查未完成"})
    # Machine checks cannot certify arbitrary overlaps, scene text or formal fonts.
    checks.append({"rule": "visual_review", "status": "needs_review",
                   "detail": "请在页面预览核对裁切、图题、字体、清晰度及学校模板"})
    status = "failed" if any(c["status"] == "failed" for c in checks) else "needs_review"
    return {**report, "checks": checks, "quality_status": status, "pages": pages}
