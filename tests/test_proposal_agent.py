# -*- coding: utf-8 -*-
from __future__ import annotations

import io
import json
import sys
from pathlib import Path
from email.message import Message

import pytest

ROOT = Path(__file__).resolve().parent.parent
SERVER = ROOT / "server"
for path in (str(SERVER), str(ROOT)):
    if path not in sys.path:
        sys.path.insert(0, path)

if __import__("importlib").util.find_spec("fastapi") is None:
    pytest.skip("server dependencies are not installed", allow_module_level=True)

from docx import Document  # noqa: E402
from pypdf import PdfWriter  # noqa: E402
from PIL import Image  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402

import app as app_module  # noqa: E402
from proposal.agent import OllamaAgent, _limit_followup_questions  # noqa: E402
from proposal import documents as document_parser  # noqa: E402
from proposal.retrieval import ProjectRetriever  # noqa: E402
from proposal.storage import Store  # noqa: E402
from proposal.tools import ProposalTools  # noqa: E402


class NoEmbeddings:
    def embed(self, texts):
        return None


def test_agent_limits_each_reply_to_one_question():
    assert _limit_followup_questions("你想研究什么？比如要解决什么问题？") == "你想研究什么？"
    assert _limit_followup_questions("你有模板吗？有的话可以上传，没有的话我用通用模板。") == (
        "你有模板吗？有的话可以上传，没有的话我用通用模板。"
    )


def test_openai_compatible_adapter_maps_tool_calls_and_keeps_key_in_header(workspace, monkeypatch):
    _client, store, _retriever, tools, _agent = workspace
    import proposal.agent as agent_module

    captured = {}

    def fake_urlopen(request, timeout=0):
        captured["request"] = request
        captured["body"] = json.loads(request.data.decode("utf-8"))
        return io.BytesIO(
            json.dumps(
                {
                    "choices": [
                        {
                            "message": {
                                "role": "assistant",
                                "content": "",
                                "tool_calls": [
                                    {
                                        "id": "call_test_1",
                                        "type": "function",
                                        "function": {"name": "search_project_materials", "arguments": "{\"query\":\"背景\"}"},
                                    }
                                ],
                            }
                        }
                    ]
                }
            ).encode()
        )

    monkeypatch.setattr(agent_module.urllib.request, "urlopen", fake_urlopen)
    remote_agent = OllamaAgent(
        store,
        tools,
        base_url="https://api.example.test/v1",
        model="example-chat",
        provider="openai_compatible",
        api_key="secret-test-key",
    )
    result = remote_agent._post(
        "/api/chat",
        {
            "model": "example-chat",
            "messages": [
                {"role": "user", "content": "查材料"},
                {
                    "role": "assistant",
                    "content": "",
                    "tool_calls": [
                        {
                            "id": "call_test_0",
                            "function": {
                                "name": "get_project_state",
                                "arguments": {"project_id": "ignored"},
                            },
                        }
                    ],
                },
                {
                    "role": "tool",
                    "tool_name": "get_project_state",
                    "tool_call_id": "call_test_0",
                    "content": "{}",
                },
            ],
            "tools": [],
            "stream": False,
        },
    )
    assert captured["request"].full_url == "https://api.example.test/v1/chat/completions"
    assert captured["request"].get_header("Authorization") == "Bearer secret-test-key"
    assert captured["body"]["messages"][-1]["tool_call_id"] == "call_test_0"
    assert captured["body"]["messages"][-1]["role"] == "tool"
    assert result["message"]["tool_calls"][0]["id"] == "call_test_1"
    assert result["message"]["tool_calls"][0]["function"]["name"] == "search_project_materials"


def test_chat_context_marks_uploaded_material_unconfirmed(workspace, tmp_path):
    _client, store, _retriever, _tools, agent = workspace
    project = store.create_project()
    file_path = tmp_path / "proposal-notes.txt"
    file_path.write_text("项目摘要内容", encoding="utf-8")
    document = store.add_document(
        project["id"], "项目摘要.txt", str(file_path), "source", "text/plain"
    )
    store.update_document_extraction(document["id"], "项目摘要内容", "ready", confirmed=False)
    captured = {}

    def fake_post(_endpoint, payload, **_kwargs):
        captured["messages"] = payload["messages"]
        return {"message": {"role": "assistant", "content": "收到，请先核对识别文字。"}}

    agent._post = fake_post
    agent.chat(project["id"], "我刚上传了项目材料。")
    system_context = captured["messages"][0]["content"]
    assert '"filename": "项目摘要.txt"' in system_context
    assert '"confirmed": false' in system_context


@pytest.fixture()
def workspace(monkeypatch, tmp_path):
    store = Store(tmp_path / "agent-data")
    retriever = ProjectRetriever(store, NoEmbeddings())
    tools = ProposalTools(store, retriever)
    agent = OllamaAgent(store, tools)
    monkeypatch.setattr(agent, "is_available", lambda: False)
    monkeypatch.setattr(
        agent,
        "status",
        lambda: {"running": False, "model_installed": False, "model": agent.model, "local": True},
    )
    monkeypatch.setattr(app_module, "store", store)
    monkeypatch.setattr(app_module, "retriever", retriever)
    monkeypatch.setattr(app_module, "tools", tools)
    monkeypatch.setattr(app_module, "agent", agent)
    with TestClient(app_module.app) as client:
        yield client, store, retriever, tools, agent


def create_project(client):
    response = client.post(
        "/api/projects",
        json={
            "title": "校园植物病害识别",
            "level": "校级",
            "discipline": "植物保护",
            "school": "示例大学",
        },
    )
    assert response.status_code == 200
    return response.json()["id"]


def test_project_persistence_and_local_session(workspace):
    client, store, *_ = workspace
    project_id = create_project(client)
    details = client.get(f"/api/projects/{project_id}")
    assert details.status_code == 200
    assert details.json()["title"] == "校园植物病害识别"
    assert client.get("/api/projects").json()["projects"][0]["id"] == project_id
    assert store.get_project(project_id)["discipline"] == "植物保护"
    assert store.get_project(project_id)["state"]["project_name"] == "校园植物病害识别"


def test_default_project_title_is_not_saved_as_a_user_fact(workspace):
    client, store, *_ = workspace
    response = client.post("/api/projects", json={"title": ""})
    assert response.status_code == 200
    project = store.get_project(response.json()["id"])
    assert project["title"] == "我的大创项目"
    assert "project_name" not in project["state"]


def test_unconfirmed_uploaded_material_is_not_retrievable(workspace):
    client, store, retriever, tools, _agent = workspace
    project_id = create_project(client)
    payload = "团队已经完成 12 株样本观察；导师建议进一步比较叶片病斑。".encode()
    response = client.post(
        f"/api/projects/{project_id}/documents?kind=source",
        files={"files": ("project-notes.txt", payload, "text/plain")},
    )
    assert response.status_code == 200
    document = response.json()["documents"][0]["document"]
    assert not document["confirmed"]
    assert retriever.search(project_id, "12 株样本")["results"] == []

    confirm = client.patch(
        f"/api/documents/{document['id']}/confirm",
        json={"confirmed": True, "corrected_text": "已完成 12 株样本观察；导师建议比较病斑。"},
    )
    assert confirm.status_code == 200
    result = tools.search_project_materials(project_id, "12 株样本", 3)
    assert result["results"]
    assert result["results"][0]["source_name"] == "project-notes.txt"
    assert result["retrieval_mode"] == "keyword_fallback"

    deleted = client.delete(f"/api/documents/{document['id']}")
    assert deleted.status_code == 200
    assert retriever.search(project_id, "12 株样本")["results"] == []


def test_mcp_text_material_requires_confirmation_before_rag(workspace):
    _client, store, retriever, tools, _agent = workspace
    project = store.create_project()
    added = tools.add_text_material(
        project["id"],
        "../../field-notes.txt",
        "观察了 12 株番茄苗，并记录病斑面积。",
    )
    assert added["confirmed"] is False
    assert added["filename"] == "field-notes.txt"
    assert retriever.search(project["id"], "12 株番茄苗")["results"] == []

    confirmed = tools.confirm_project_material(project["id"], added["document_id"])
    assert confirmed["success"] is True
    assert confirmed["indexed_chunks"] > 0
    assert retriever.search(project["id"], "12 株番茄苗")["results"]


def test_local_ocr_extracts_photo_and_scanned_pdf(monkeypatch):
    class FakeOCR:
        def __call__(self, _image):
            return type("OCRResult", (), {"txts": ["观察了 12 株番茄苗"], "scores": [0.61]})()

    monkeypatch.setattr(document_parser, "_ocr_engine", lambda: FakeOCR())
    image = Image.new("RGB", (600, 200), "white")
    image_bytes = io.BytesIO()
    image.save(image_bytes, format="PNG")
    image_data = image_bytes.getvalue()

    photo = document_parser.extract_file("experiment.png", image_data)
    assert photo.status == "ready"
    assert "12 株" in photo.text
    assert photo.low_confidence_count == 1

    pdf_bytes = io.BytesIO()
    image.save(pdf_bytes, format="PDF")
    scanned = document_parser.extract_file("scan.pdf", pdf_bytes.getvalue())
    assert scanned.status == "ready"
    assert scanned.source_pages[0][0] == "第 1 页"
    assert "番茄苗" in scanned.text


def test_outline_confirmation_section_edit_and_export(workspace):
    client, store, _retriever, _tools, agent = workspace
    project_id = create_project(client)
    agent.propose_outline = lambda _pid: [
        {"id": "background", "title": "项目背景与研究问题"},
        {"id": "research_content", "title": "研究内容与方法"},
    ]
    proposed = client.post(f"/api/projects/{project_id}/outline/generate")
    assert proposed.status_code == 200
    assert proposed.json()["confirmed"] is False
    assert store.get_project(project_id)["state"].get("outline_confirmed") is False

    rejected_generation = client.post(f"/api/projects/{project_id}/sections/background/generate")
    assert rejected_generation.status_code == 400

    confirmed = client.put(
        f"/api/projects/{project_id}/outline",
        json={"outline": proposed.json()["outline"]},
    )
    assert confirmed.status_code == 200
    assert store.get_project(project_id)["state"]["outline_confirmed"] is True
    saved = client.put(
        f"/api/projects/{project_id}/sections/background",
        json={"content": "校园植物病害会影响苗木存活，拟通过图像识别辅助早期发现。"},
    )
    assert saved.status_code == 200

    export = client.post(f"/api/projects/{project_id}/export")
    assert export.status_code == 200
    assert export.json()["draft"] is True
    assert "download_url" in export.json()
    download = client.get(export.json()["download_url"])
    assert download.status_code == 200
    generated = Document(io.BytesIO(download.content))
    text = "\n".join(p.text for p in generated.paragraphs)
    assert "校园植物病害" in text
    assert "大学生创新创业训练计划" in text
    assert "待补充" in text
    assert "项目实践照片" in text


def test_non_innovation_skill_uses_skill_title_and_renders_markdown_tables(workspace):
    client, store, _retriever, _tools, _agent = workspace
    project = client.post(
        "/api/projects",
        json={"skill_id": "motivation-scholarship", "has_template": False},
    ).json()
    project_id = project["id"]
    client.put(
        f"/api/projects/{project_id}/outline",
        json={"outline": [{"id": "family", "title": "家庭经济情况"}]},
    )
    client.put(
        f"/api/projects/{project_id}/sections/family",
        json={
            "content": "| 成员 | 年收入 |\n|---|---|\n| 父亲 | 待补充 |\n\n"
            "学校认定等级：待补充。"
        },
    )
    result = client.post(f"/api/projects/{project_id}/export")
    assert result.status_code == 200
    document = Document(io.BytesIO(client.get(result.json()["download_url"]).content))
    paragraphs = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "国家励志奖学金" in paragraphs
    assert "大学生创新创业训练计划" not in paragraphs
    assert any(table.cell(1, 0).text == "父亲" for table in document.tables)


def test_missing_facts_produce_placeholders_not_invented_section_or_diagram(workspace):
    _client, store, _retriever, tools, agent = workspace
    project = store.create_project(
        "验收用项目名称",
        "校级",
        "环境科学",
        "示例大学",
    )
    store.update_project(
        project["id"],
        state={"project_name": "验收用项目名称", "outline_confirmed": True},
        outline=[
            {"id": "project_info", "title": "项目基本信息"},
            {"id": "background", "title": "项目背景与研究问题"},
        ],
    )

    def fail_if_model_called(*_args, **_kwargs):
        raise AssertionError("缺少事实时不调用模型扩写项目事实")

    agent._post = fail_if_model_called
    basic = agent.generate_section(project["id"], "project_info")
    assert "验收用项目名称" in basic["content"]
    assert "待补充" in basic["content"]
    background = agent.generate_section(project["id"], "background")
    assert "待补充" in background["content"]
    diagram = agent.propose_diagram(project["id"], "技术路线图")
    assert all("待补充" in node for node in diagram["nodes"])


def test_word_template_is_used_and_labeled_cells_are_filled(workspace, tmp_path):
    client, store, _retriever, _tools, _agent = workspace
    project_id = create_project(client)
    template = Document()
    template.add_heading("学校大创申报模板", 0)
    table = template.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "项目名称"
    table.cell(1, 0).text = "立项依据"
    template_bytes = io.BytesIO()
    template.save(template_bytes)
    template_bytes.seek(0)
    upload = client.post(
        f"/api/projects/{project_id}/documents?kind=template",
        files={"files": ("school-template.docx", template_bytes.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert upload.status_code == 200
    template_doc = upload.json()["documents"][0]["document"]
    assert template_doc["confirmed"] is True
    assert store.update_project(
        project_id,
        state={"project_name": "校园植物病害识别"},
        outline=[{"id": "background", "title": "立项依据"}],
        sections={"background": "拟利用图像方法识别植物叶片病害。"},
    )

    export = client.post(f"/api/projects/{project_id}/export")
    assert export.status_code == 200
    body = export.json()
    assert body["used_original_template"] is True
    assert body["template_sections_filled"] >= 1
    doc = Document(io.BytesIO(client.get(body["download_url"]).content))
    assert doc.tables[0].cell(0, 1).text == "校园植物病害识别"
    assert "拟利用图像方法" in doc.tables[0].cell(1, 1).text


def test_word_template_paragraph_sections_are_filled_in_place(workspace, tmp_path):
    client, store, _retriever, _tools, _agent = workspace
    project_id = create_project(client)
    template = Document()
    template.add_heading("学校大创申请书", 0)
    template.add_heading("项目背景与研究问题", level=2)
    template.add_paragraph("（请填写项目背景）")
    file_data = io.BytesIO()
    template.save(file_data)
    uploaded = client.post(
        f"/api/projects/{project_id}/documents?kind=template",
        files={"files": ("paragraph-template.docx", file_data.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    ).json()["documents"][0]["document"]
    store.set_template_document(project_id, uploaded["id"])
    store.update_project(
        project_id,
        outline=[{"id": "background", "title": "项目背景与研究问题"}],
        sections={"background": "本项目拟观察土壤水分变化与植物病斑表现之间的关系。"},
    )
    response = client.post(f"/api/projects/{project_id}/export")
    assert response.status_code == 200
    assert response.json()["template_paragraphs_filled"] == 1
    body = client.get(response.json()["download_url"])
    document = Document(io.BytesIO(body.content))
    assert document.paragraphs[2].text == "本项目拟观察土壤水分变化与植物病斑表现之间的关系。"


def test_pdf_template_headings_are_used_as_reconstruction_order(workspace, tmp_path):
    client, store, _retriever, _tools, _agent = workspace
    project_id = create_project(client)
    pdf_path = store.files_dir / project_id / "school-template.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    with pdf_path.open("wb") as handle:
        writer.write(handle)
    document = store.add_document(
        project_id,
        "school-template.pdf",
        str(pdf_path),
        "template",
        "application/pdf",
    )
    store.update_document_extraction(
        document["id"],
        "[第 1 页]\n一、项目背景与研究问题\n二、研究方法\n三、进度安排",
        "ready",
        confirmed=True,
    )
    store.set_template_document(project_id, document["id"])
    store.update_project(
        project_id,
        outline=[
            {"id": "background", "title": "项目背景与研究问题"},
            {"id": "method", "title": "研究方法"},
            {"id": "schedule", "title": "进度安排"},
        ],
        sections={"background": "待补充", "method": "待补充", "schedule": "待补充"},
    )
    response = client.post(f"/api/projects/{project_id}/export")
    assert response.status_code == 200
    assert response.json()["pdf_template_outline_applied"] is True
    assert any("PDF 模板" in warning for warning in response.json()["warnings"])


def test_official_source_import_requires_review_before_rag(workspace, monkeypatch):
    client, store, _retriever, tools, _agent = workspace
    project_id = create_project(client)
    monkeypatch.setattr(
        tools,
        "fetch_official_notice",
        lambda _url: {
            "success": True,
            "title": "2026 大创申报通知",
            "url": "https://example.edu.cn/notice/2026",
            "content": "项目申请截止日期为 2026 年 5 月 10 日。",
            "published_at": "2026-04-01",
            "retrieved_at_utc": "2026-10-08T00:00:00+00:00",
            "source_domain": "example.edu.cn",
            "warning": "请核对通知适用范围。",
        },
    )
    imported = client.post(
        f"/api/projects/{project_id}/official/import",
        json={"url": "https://example.edu.cn/notice/2026"},
    )
    assert imported.status_code == 200
    document = imported.json()["document"]
    assert document["kind"] == "official"
    assert document["confirmed"] is False
    assert store.get_project(project_id)["state"].get("official_sources", []) == []

    confirmed = client.patch(
        f"/api/documents/{document['id']}/confirm",
        json={"confirmed": True, "corrected_text": document["extracted_text"]},
    )
    assert confirmed.status_code == 200
    project_state = store.get_project(project_id)["state"]
    assert project_state["official_sources"][0]["url"] == "https://example.edu.cn/notice/2026"
    assert store.chunks_for_project(project_id)


def test_fact_tool_requires_verifiable_quote(workspace):
    client, store, _retriever, tools, _agent = workspace
    project_id = create_project(client)
    store.add_message(project_id, "user", "我们已经观察了 12 株番茄苗。")
    pending = tools.save_project_fact(
        project_id,
        "foundation",
        "我们已经观察了 12 株番茄苗。",
        "我们已经观察了 12 株番茄苗。",
    )
    assert pending["pending_confirmation"] is True
    assert "foundation" not in store.get_project(project_id)["state"]
    store.add_message(project_id, "user", "确认")
    rejected = tools.save_project_fact(
        project_id, "problem", "已经观察 1200 株", "我们已经观察了 12 株番茄苗。"
    )
    assert "逐字一致" in rejected["error"]
    accepted = tools.save_project_fact(
        project_id,
        "foundation",
        "我们已经观察了 12 株番茄苗。",
        "我们已经观察了 12 株番茄苗。",
    )
    assert accepted["success"] is True
    project = store.get_project(project_id)
    assert project["state"]["foundation"] == "我们已经观察了 12 株番茄苗。"
    assert project["state"]["fact_sources"]["foundation"] == "我们已经观察了 12 株番茄苗。"


def test_fact_candidates_are_unsaved_until_user_confirms(workspace):
    client, store, _retriever, _tools, agent = workspace
    project_id = create_project(client)
    store.add_message(project_id, "user", "我们已经观察了 12 株番茄苗。")
    agent._post = lambda *_args, **_kwargs: {
        "message": {
            "content": json.dumps(
                {
                    "facts": [
                        {
                            "key": "foundation",
                            "value": "我们已经观察了 12 株番茄苗。",
                            "source_quote": "我们已经观察了 12 株番茄苗。",
                        }
                    ]
                },
                ensure_ascii=False,
            )
        }
    }
    proposed = client.post(f"/api/projects/{project_id}/facts/propose")
    assert proposed.status_code == 200
    candidate = proposed.json()["candidates"][0]
    assert proposed.json()["saved"] is False
    assert "foundation" not in store.get_project(project_id)["state"]

    confirmed = client.post(f"/api/projects/{project_id}/facts/confirm", json=candidate)
    assert confirmed.status_code == 200
    assert confirmed.json()["confirmed"] is True
    assert store.get_project(project_id)["state"]["foundation"] == "我们已经观察了 12 株番茄苗。"


def test_validation_flags_conflicting_budget_values_from_confirmed_docs(workspace):
    client, store, _retriever, tools, _agent = workspace
    project_id = create_project(client)
    for name, text in [
        ("project-brief.txt", "[行 1] 项目总经费：5000元。"),
        ("budget-sheet.txt", "[行 1] 项目总经费：8000元。"),
    ]:
        file_path = store.files_dir / project_id / name
        file_path.write_text(text, encoding="utf-8")
        document = store.add_document(project_id, name, str(file_path), "source", "text/plain")
        store.update_document_extraction(document["id"], text, "ready", confirmed=True)
    result = tools.validate_project(project_id)
    assert result["passed"] is False
    assert any("多个预算金额" in warning for warning in result["consistency_warnings"])


def test_validation_warns_about_unmatched_source_citations(workspace):
    _client, store, _retriever, tools, _agent = workspace
    project = store.create_project()
    store.update_project(
        project["id"],
        outline=[{"id": "background", "title": "背景"}],
        sections={"background": "研究基础来自已有材料 [来源：不存在的记录.txt，段落 2]。"},
    )
    result = tools.validate_project(project["id"])
    assert result["passed"] is False
    assert result["source_citation_warnings"]


def test_ollama_tool_call_loop_uses_shared_tool_registry(workspace):
    _client, store, _retriever, tools, agent = workspace
    project = store.create_project()
    store.add_message(project["id"], "user", "我已经观察了 12 株番茄苗。")
    responses = iter(
        [
            {
                "message": {
                    "role": "assistant",
                    "tool_calls": [
                        {
                            "function": {
                                "name": "save_project_fact",
                                "arguments": {
                                    "key": "foundation",
                                    "value": "我已经观察了 12 株番茄苗。",
                                    "source_quote": "我已经观察了 12 株番茄苗。",
                                },
                            }
                        }
                    ],
                }
            },
            {
                "message": {
                    "role": "assistant",
                    "content": "我理解你已经观察了 12 株番茄苗。是否确认保存这个基础信息？",
                }
            },
            {
                "message": {
                    "role": "assistant",
                    "tool_calls": [
                        {
                            "function": {
                                "name": "save_project_fact",
                                "arguments": {
                                    "key": "foundation",
                                    "value": "我已经观察了 12 株番茄苗。",
                                    "source_quote": "我已经观察了 12 株番茄苗。",
                                },
                            }
                        }
                    ],
                }
            },
            {"message": {"role": "assistant", "content": "已按你确认的信息保存。"}},
        ]
    )
    agent._post = lambda *_args, **_kwargs: next(responses)
    initial = agent.chat(project["id"], "我已经观察了 12 株番茄苗。")
    assert initial["tool_calls"] == ["save_project_fact"]
    assert "foundation" not in store.get_project(project["id"])["state"]
    result = agent.chat(project["id"], "确认")
    assert result["tool_calls"] == ["save_project_fact"]
    assert store.get_project(project["id"])["state"]["foundation"] == "我已经观察了 12 株番茄苗。"


def test_bare_confirmation_does_not_answer_a_multi_option_question(workspace):
    _client, store, _retriever, tools, agent = workspace
    project = store.create_project()
    store.add_message(project["id"], "assistant", "番茄苗是在实验条件下还是自然环境中观察的？")

    def should_not_call_model(*_args, **_kwargs):
        raise AssertionError("ambiguous bare confirmation should be handled deterministically")

    agent._post = should_not_call_model
    result = agent.chat(project["id"], "确认")
    assert result["tool_calls"] == []
    assert "不能替你猜" in result["message"]
    assert store.get_project(project["id"])["state"] == {}


def test_diagram_save_sanitizes_nodes_and_returns_render_status(workspace):
    client, _store, _retriever, _tools, _agent = workspace
    project_id = create_project(client)
    response = client.post(
        f"/api/projects/{project_id}/diagrams",
        json={"title": "技术路线图", "nodes": ["问题定义", "采集数据", "验证方案"]},
    )
    assert response.status_code == 200
    diagram = response.json()["diagram"]
    assert len(diagram["nodes"]) == 3
    assert diagram["title"] == "技术路线图"


def test_official_fetch_rejects_untrusted_or_insecure_urls(workspace):
    _client, _store, _retriever, tools, _agent = workspace
    assert "HTTPS" in tools.fetch_official_notice("http://example.edu.cn/notice")["error"]
    assert "只允许" in tools.fetch_official_notice("https://example.com/notice")["error"]
    assert "私有网络" in tools.fetch_official_notice("https://127.0.0.1/notice")["error"]


def test_official_page_fetch_extracts_title_date_and_visible_text(workspace, monkeypatch):
    _client, _store, _retriever, tools, _agent = workspace
    import proposal.tools as tools_module

    monkeypatch.setattr(
        tools_module.socket,
        "getaddrinfo",
        lambda *_args, **_kwargs: [(2, 1, 6, "", ("93.184.216.34", 443))],
    )

    class FakeResponse:
        def __init__(self):
            self.headers = Message()
            self.headers["Content-Type"] = "text/html; charset=utf-8"
            self.body = (
                '<html><head><title>2026 大创申报通知</title>'
                '<meta name="publishdate" content="2026-04-01"></head>'
                '<body><script>ignore this</script><h1>申报要求</h1>'
                '<p>项目申请截止日期为 2026 年 5 月 10 日。</p></body></html>'
            ).encode()

        def __enter__(self):
            return self

        def __exit__(self, *_args):
            return False

        def read(self, limit):
            return self.body[:limit]

        def geturl(self):
            return "https://example.edu.cn/notice/2026"

    class FakeOpener:
        def open(self, *_args, **_kwargs):
            return FakeResponse()

    monkeypatch.setattr(tools_module.urllib.request, "build_opener", lambda *_args: FakeOpener())
    result = tools.fetch_official_notice("https://example.edu.cn/notice/2026")
    assert result["success"] is True
    assert result["title"] == "2026 大创申报通知"
    assert result["published_at"] == "2026-04-01"
    assert "2026 年 5 月 10 日" in result["content"]
    assert "ignore this" not in result["content"]


def test_configured_searxng_returns_only_official_domains(workspace, monkeypatch):
    _client, _store, _retriever, tools, _agent = workspace
    import proposal.tools as tools_module

    monkeypatch.setenv("SEARXNG_URL", "https://search.example/")

    class FakeResponse:
        def __enter__(self):
            return self

        def __exit__(self, *_args):
            return False

        def read(self):
            return json.dumps(
                {
                    "results": [
                        {"title": "校级大创通知", "url": "https://university.edu.cn/notice", "content": "校内规则"},
                        {"title": "论坛转载", "url": "https://forum.example.com/post", "content": "非官方"},
                    ]
                }
            ).encode()

    monkeypatch.setattr(tools_module.urllib.request, "urlopen", lambda *_args, **_kwargs: FakeResponse())
    result = tools.search_official_notices("大创申报", "示例大学", max_results=5)
    assert result["source"] == "configured SearXNG"
    assert [item["url"] for item in result["results"]] == ["https://university.edu.cn/notice"]


def test_optional_review_uses_saved_draft_and_is_explicit(workspace):
    _client, store, _retriever, tools, _agent = workspace
    project = store.create_project()
    store.update_project(
        project["id"],
        outline=[{"id": "background", "title": "项目背景"}],
        sections={"background": "项目拟研究校园番茄苗病斑与土壤水分的关系，具体数据待补充。"},
    )
    result = tools.simulate_review(project["id"])
    assert result["success"] is True
    assert result["report"]["max_score"] == 100
    assert "不代表真实评审结果" in result["notice"]
