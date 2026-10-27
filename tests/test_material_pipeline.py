"""Pilot tests: immutable inputs, dependencies, figures and durable execution.

No test calls a real text/image model or modifies a production database.
"""
import copy
import io
import json
import sys
import time
import zipfile
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "server"))
from proposal.materials import (
    RULES, TRACKING, MaterialConflict, context_values, material_view, requirements,
)
from proposal.material_figures import FigureSpec, render_python
from proposal.material_export import blocks, export_material, check_pages
from proposal.material_jobs import MaterialJobs, checked_bundle, job_id
from proposal.storage import Store
from proposal.tools import ProposalTools
from proposal.agent import OllamaAgent
from proposal.workflow import REVIEW_WEIGHTS


@pytest.fixture
def pilot(monkeypatch, tmp_path):
    monkeypatch.setenv("MATERIAL_PIPELINE_ENABLED", "true")
    store = Store(tmp_path)
    project = store.create_project(title="真实项目", school="测试大学")
    project = store.update_project(project["id"], state={
        "skill_id": "innovation-research", "outline_confirmed": True,
        "method": "滚动规划测试", "budget": "20000",
        "confirmed_facts": {"advisor_name": "真实教师"},
    }, outline=[{"id": k, "title": v} for k, v in [
        ("abstract", "摘要"), ("research_route", "技术路线"), ("budget", "预算"), ("foundation", "团队基础")]],
        sections={"abstract": "原摘要", "research_route": "原方法", "budget": "原预算", "foundation": "原团队"},
    )
    return store, project


def test_context_plans_are_not_achievements_and_unknown_types_unchanged(pilot):
    store, project = pilot
    facts, plans = context_values(project)
    assert plans["method"] == "滚动规划测试"
    assert facts["advisor_name"] == "真实教师"
    assert "method" not in facts
    view = material_view(project)
    assert view["facts"]["advisor_name"]["verification"] == "user_provided_not_externally_verified"
    other = store.update_project(project["id"], state={"skill_id": "national-scholarship"})
    assert "material" not in other
    assert material_view(other)["pipeline"] == "legacy"


def rule(value, kind="official", **changes):
    return {"key": "font", "value": value, "source_kind": kind, "source": "通知",
            "scope": "测试大学 2026 大创", "confirmed": True, "applicable": True, **changes}


def test_requirements_priority_ties_and_unconfirmed(pilot):
    _, project = pilot
    project["state"][RULES] = [rule("宋体"), rule("黑体", "example")]
    resolved = requirements(project)
    assert resolved["effective"]["font"]["value"] == "宋体"
    assert resolved["conflicts"][0]["status"] == "priority_applied"
    project["state"][RULES].append(rule("楷体"))
    assert "font" not in requirements(project)["effective"]
    assert requirements(project)["conflicts"][0]["status"] == "needs_confirmation"
    project["state"][RULES] = [rule("宋体", scope=""), rule("黑体", "user", confirmed=False)]
    assert not requirements(project)["effective"]
    assert len(requirements(project)["pending"]) == 2


def test_dependencies_acknowledgement_and_noop(pilot):
    store, p = pilot
    original = copy.deepcopy(p["sections"])
    p = store.update_project(p["id"], state={**p["state"], "method": "新方法"})
    view = material_view(p)
    assert "section:abstract" in view["stale"]
    assert "section:foundation" not in view["stale"]
    assert p["sections"] == original
    snapshot = view["snapshot"]
    p = store.update_project(p["id"], state={**p["state"],
        "_material_ack": {"snapshot": snapshot, "nodes": list(view["stale"])}})
    assert not material_view(p)["stale"]
    p = store.update_project(p["id"], state=p["state"])
    assert not material_view(p)["stale"]
    p = store.update_project(p["id"], state={**p["state"], "method": "再次改变"})
    assert material_view(p)["stale"]
    # Old acknowledgement cannot clear new changes.
    p = store.update_project(p["id"], state={**p["state"],
        "_material_ack": {"snapshot": snapshot, "nodes": ["section:abstract"]}})
    assert "section:abstract" in material_view(p)["stale"]


def test_source_change_marks_outputs_and_cas_preserves_edit(pilot):
    store, p = pilot
    file = store.files_dir / p["id"] / "source.txt"
    file.write_text("真实依据", encoding="utf-8")
    doc = store.add_document(p["id"], "source.txt", str(file), "source", "text/plain")
    store.update_document_extraction(doc["id"], "真实依据", "ready", confirmed=True, confirmed_by_user=True)
    latest = store.get_project(p["id"])
    assert material_view(latest)["stale"]
    with pytest.raises(MaterialConflict):
        store.update_project(p["id"], sections={"abstract": "覆盖"}, expected_revision=p["material_revision"])
    assert store.get_project(p["id"])["sections"]["abstract"] == "原摘要"


def payload(store, p, *, paid=False):
    return {"job_id": job_id(), "owner_id": "local_user", "project": copy.deepcopy(p),
            "base_revision": p["material_revision"], "input_snapshot": material_view(p)["snapshot"],
            "documents": [], "messages": [], "chunks": [], "paid": paid, "platform": False,
            "downloads": [], "config": {"provider": "ollama"}}


def finish_section(queue, job, content="生成正文"):
    running = queue.claim_local()
    assert running["id"] == job["id"]
    queue.worker_local(running, "finish", status="completed", result={"content": content})
    return queue.call("raw", project_id=job["project_id"], id=job["id"])


def test_idempotency_persistence_and_apply_preserves_manual_clear(pilot):
    store, p = pilot
    queue = MaterialJobs(store)
    job = queue.submit(p["id"], "section", "abstract", payload(store, p), "same-click")
    assert queue.submit(p["id"], "section", "abstract", payload(store,p), "same-click")["id"] == job["id"]
    assert MaterialJobs(Store(store.data_dir)).call("get", project_id=p["id"], id=job["id"])["id"] == job["id"]
    raw = finish_section(queue, job)
    store.update_project(p["id"], sections={**p["sections"], "abstract": ""})
    result = queue.apply(raw, {"sections": {**p["sections"], "abstract": "生成正文"}})
    assert result["conflict_status"] == "user_edit_preserved"
    assert store.get_project(p["id"])["sections"]["abstract"] == ""
    assert store.list_section_versions(p["id"], "abstract")[0]["content"] == "生成正文"
    assert queue.apply(raw, {}) == result
    assert len(store.list_section_versions(p["id"], "abstract")) == 1


def test_atomic_apply_and_cross_project_rejection(pilot):
    store, p = pilot
    queue = MaterialJobs(store)
    job = queue.submit(p["id"], "section", "abstract", payload(store, p), "normal")
    raw = finish_section(queue, job)
    result = queue.apply(raw, {"sections": {**p["sections"], "abstract": "生成正文"}})
    assert result["applied"]
    assert store.get_project(p["id"])["sections"]["abstract"] == "生成正文"
    another = store.create_project()
    assert queue.call("get", project_id=another["id"], id=job["id"]) is None
    assert queue.call("cancel", project_id=another["id"], id=job["id"]) is None


def test_expired_paid_lease_never_auto_replays(pilot):
    store, p = pilot
    queue = MaterialJobs(store)
    job = queue.submit(p["id"], "section", "abstract", payload(store,p,paid=True), "paid")
    running = queue.claim_local()
    with store.connect() as db:
        running["lease_until"] = time.time() - 1
        db.execute("UPDATE material_jobs SET body=? WHERE id=?", (json.dumps(running), job["id"]))
    assert queue.claim_local() is None
    assert queue.call("get", project_id=p["id"], id=job["id"])["status"] == "needs_review"


def test_cancel_blocks_checkpoint_and_terminal_write(pilot):
    store, p = pilot
    queue = MaterialJobs(store)
    job = queue.submit(p["id"], "export", "", payload(store,p), "cancel")
    running = queue.claim_local()
    queue.call("cancel", project_id=p["id"], id=job["id"])
    with pytest.raises(MaterialConflict):
        queue.worker_local(running, "finish", status="completed", result={})


def test_real_generation_path_does_not_write_before_apply(pilot, monkeypatch):
    store, p = pilot
    tools = ProposalTools(store)
    monkeypatch.setattr(tools, "search_project_materials", lambda *a,**k: {"results": [{"text":"已确认依据"}]})
    agent = OllamaAgent(store, tools)
    prompts = []
    review = {"overall_score":85.0,"passed":True,
              "dimensions":[{"name":n,"score":85.0,"weight":w,"comment":"检查"} for n,w in REVIEW_WEIGHTS.items()],
              "strengths":[],"weaknesses":[],"actionable_revisions":[],"user_action_required":[],"blocking_issues":[]}
    def fake_post(endpoint, request, **kwargs):
        prompts.append(json.dumps(request["messages"], ensure_ascii=False))
        return {"message":{"content": "有依据的草稿" if len(prompts)==1 else json.dumps(review)}}
    monkeypatch.setattr(agent, "_post", fake_post)
    result = agent.generate_section(p["id"], "research_route", persist=False)
    assert "滚动规划测试" in prompts[0]
    assert "真实教师" in prompts[0]
    assert result["content"] == "有依据的草稿"
    assert store.get_project(p["id"])["sections"]["research_route"] == "原方法"


@pytest.mark.parametrize("overrides", [
    {"renderer":"python","labels":["A"],"values":[float("nan")],"source":"来源"},
    {"renderer":"python","labels":["A"],"values":[1],"source":""},
    {"renderer":"svg","nodes":[""]},
    {"renderer":"svg","nodes":["正常"],"script":"任意代码"},
])
def test_figures_reject_unbounded_or_untrusted_code(overrides):
    with pytest.raises(ValueError):
        FigureSpec(id="fig",section_id="abstract",title="测试图",**overrides)


def test_python_planned_data_and_docx_chapter_placement(pilot, tmp_path):
    store, p = pilot
    folder = tmp_path / "figure"; folder.mkdir()
    spec = FigureSpec(id="fig",section_id="research_route",title="计划指标",
        renderer="python",labels=["阶段一","阶段二"],values=[2,3],source="用户计划")
    image = render_python(spec,folder)
    assert image.is_file()
    p = store.update_project(p["id"], diagrams=[{**spec.model_dump(),"preview":str(image)}])
    output = tmp_path / "output"
    report = export_material(store,p["id"],output)
    from docx import Document
    document = Document(output / "proposal.docx")
    paragraphs = [x.text for x in document.paragraphs]
    figure_at = next(i for i,x in enumerate(paragraphs) if x.startswith("图 1"))
    assert next(i for i,x in enumerate(paragraphs) if "技术路线" in x) < figure_at
    assert figure_at < next(i for i,x in enumerate(paragraphs) if "、预算" in x)
    assert len(document.inline_shapes) == 1
    assert report["layout_profile"]["left_cm"] == 3
    assert blocks("普通段落\n[[figure:fig]]")[1]["type"] == "figure"


def test_template_placeholder_is_filled_in_place_once_with_structured_blocks(pilot, tmp_path):
    store, p = pilot
    from docx import Document
    template = tmp_path / "school-template.docx"
    doc = Document()
    doc.add_paragraph("学校模板说明：保留原页眉和栏目")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "立项依据"
    table.cell(0, 1).text = "请填写"
    table.cell(1, 0).text = "技术路线"
    table.cell(1, 1).text = "{{section:research_route}}"
    doc.add_paragraph("固定尾注：学校落款")
    doc.save(template)
    source = store.add_document(p["id"], "school-template.docx", str(template), "template",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    store.set_template_document(p["id"], source["id"])
    store.update_project(p["id"], sections={**p["sections"], "research_route":
        "第一阶段开展资料核对。\n|检查项|状态|\n|---|---|\n|事实|待核对|"})
    report = export_material(store, p["id"], tmp_path / "template-out")
    filled = Document(tmp_path / "template-out" / "proposal.docx")
    def cell_text(cell):
        return cell.text + "\n" + "\n".join(cell_text(c) for table in cell.tables for row in table.rows for c in row.cells)
    all_text = "\n".join(p.text for p in filled.paragraphs) + "\n" + "\n".join(
        cell_text(c) for table in filled.tables for row in table.rows for c in row.cells)
    assert "学校模板说明" in all_text and "学校落款" in all_text
    assert "第一阶段开展资料核对" in all_text and "待核对" in all_text
    assert "{{section:research_route}}" not in all_text
    assert "第一阶段开展资料核对" not in "\n".join(p.text for p in filled.paragraphs)
    assert "原学校模板" not in " ".join(report["warnings"])


def test_bundle_rejects_zip_slip_and_bombs(tmp_path):
    data = io.BytesIO()
    with zipfile.ZipFile(data, "w") as archive:
        archive.writestr("../escape.png", b"bad")
    with pytest.raises(ValueError):
        checked_bundle(data.getvalue(), tmp_path)


def test_worker_export_without_model_calls(pilot, monkeypatch):
    store, p = pilot
    import material_worker
    def forbidden(*args, **kwargs):
        pytest.fail("Export must not call paid models")
    monkeypatch.setattr(OllamaAgent, "_post", forbidden)
    queue = MaterialJobs(store)
    job = queue.submit(p["id"], "export", "", payload(store,p), "export")
    assert material_worker.run_one(local_store=store)
    done = queue.call("get", project_id=p["id"], id=job["id"])
    assert done["status"] == "completed", done
    report = done["result"]["report"]
    assert report["quality_status"] in {"needs_review","not_checked","failed"}
    assert (store.files_dir / p["id"] / "_material" / job["id"] / "proposal.docx").is_file()
