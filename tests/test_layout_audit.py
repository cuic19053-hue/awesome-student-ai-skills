#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""排版硬指标审计器 + docx_common 防拆构件的回归测试。

这两件事必须绑在 CI 上：审计器一旦漏检，所有 skill 会静默回到「表格跨页、
图注分家、图表零引用」的老样子；构件一旦丢保护，同样的问题会重新出现在
用户文档里。所以这里既测「审计能抓到问题」，也测「构件真的加了保护」。
"""

from __future__ import annotations

import sys
from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "utils"))

from utils import docx_common as dc  # noqa: E402
from utils import layout_audit as la  # noqa: E402


# ---------------------------------------------------------------- 构件保护

def _row_cant_split(row) -> bool:
    trPr = row._tr.find(qn("w:trPr"))
    return trPr is not None and trPr.find(qn("w:cantSplit")) is not None


def test_add_table_hardens_rows():
    doc = Document()
    t = dc.add_table(doc, ["指标", "目标值"], [("成功率", "94.3%"), ("耗时", "168ms")])
    assert all(_row_cant_split(r) for r in t.rows)


def test_long_table_repeats_header():
    doc = Document()
    rows = [(f"行{i}", i) for i in range(10)]
    t = dc.add_table(doc, ["项目", "值"], rows)
    trPr = t.rows[0]._tr.find(qn("w:trPr"))
    assert trPr is not None and trPr.find(qn("w:tblHeader")) is not None


def test_short_table_keeps_together():
    doc = Document()
    t = dc.add_table(doc, ["项目", "值"], [("A", 1), ("B", 2)])
    for row in t.rows[:-1]:
        for cell in row.cells:
            assert cell.paragraphs[0].paragraph_format.keep_with_next


def test_page_number_field_structure():
    """域必须拆成 begin/instr/separate/结果/end 五个 run。

    早期实现把四个节点塞进同一个 run，Word 打开后页码显示为空。
    """
    doc = Document()
    dc.add_page_number(doc)
    footer_xml = doc.sections[0].footer.paragraphs[0]._p.xml
    assert "PAGE" in footer_xml and "NUMPAGES" in footer_xml
    assert doc.sections[0].different_first_page_header_footer is True
    # separate 与 end 之间必须有文本 run，否则显示结果是空白
    p = doc.sections[0].footer.paragraphs[0]
    texts = [r.text for r in p.runs]
    assert any(t.strip() for t in texts), "域结果 run 缺失，页码会显示为空"


def test_review_boxes_one_row_per_box():
    doc = Document()
    t = dc.add_review_boxes(doc, ["指导教师意见", "学院审核意见", "学校评审意见"])
    assert len(t.rows) == 3                    # 一栏一行
    assert all(_row_cant_split(r) for r in t.rows)
    assert len(t.columns) == 1                 # 一列，栏内不再拆行


# ---------------------------------------------------------------- 审计器

def _build_good(tmp: Path) -> Path:
    """结构完整的最小文档：页码 + 表题 + 表格保护 + 图注绑定 + 图表引用。"""
    out = tmp / "good.docx"
    doc = dc.create_docx(out, title="测试文档")
    dc.add_heading1(doc, "一、基本情况")
    dc.add_paragraph(doc, "本项目的核心指标见表 1。")
    t = dc.add_table(doc, ["指标", "目标"], [("成功率", "94.3%")])
    dc.add_paragraph(doc, "详细数据见表 1 的逐项说明。")
    dc.add_heading1(doc, "二、说明")
    dc.add_paragraph(doc, "以下内容不含图表。")
    dc.add_page_number(doc)
    doc.save(str(out))
    return out


def test_audit_passes_clean_doc(tmp_path):
    p = _build_good(tmp_path)
    rep = la.audit(p, render=False)
    assert not rep.fails, [f"{i.rule}: {i.detail}" for i in rep.fails]
    assert rep.score >= 90


def test_audit_catches_missing_page_number(tmp_path):
    doc = dc.create_docx("无页码")
    dc.add_heading1(doc, "一、基本情况")
    dc.add_paragraph(doc, "正文。")
    out = tmp_path / "nopage.docx"
    doc.save(str(out))
    rep = la.audit(out, render=False)
    assert any(i.rule == "L06" for i in rep.fails)


def test_audit_catches_orphan_caption(tmp_path):
    """有图注、但正文从不引用编号 —— L04 必须报出来。"""
    doc = dc.create_docx("无引用")
    dc.add_heading1(doc, "一、基本情况")
    dc.add_paragraph(doc, "下面这张图说明问题。")
    cap = doc.add_paragraph("图 1\u3000整体流程")
    dc.add_page_number(doc)
    out = tmp_path / "nocite.docx"
    doc.save(str(out))
    rep = la.audit(out, render=False)
    assert any(i.rule == "L04" for i in rep.fails)


def test_audit_ignores_word_internal_false_positive(tmp_path):
    """「累计发表 14 篇」不能被判成「引用了表 14」。"""
    doc = dc.create_docx("假阳性")
    dc.add_heading1(doc, "一、基本情况")
    dc.add_paragraph(doc, "在校报累计发表 14 篇科普短文。")
    dc.add_page_number(doc)
    out = tmp_path / "fp.docx"
    doc.save(str(out))
    rep = la.audit(out, render=False)
    assert not any("表 [14]" in i.detail or "表 14" in i.detail for i in rep.issues)


def test_audit_placeholder_residue(tmp_path):
    doc = dc.create_docx("占位")
    dc.add_heading1(doc, "一、基本情况")
    dc.add_paragraph(doc, "此处待补充具体数据。")
    dc.add_page_number(doc)
    out = tmp_path / "todo.docx"
    doc.save(str(out))
    rep = la.audit(out, render=False)
    assert any(i.rule == "L12" for i in rep.fails)


if __name__ == "__main__":
    sys.exit(pytest.main([__file__, "-q"]))
