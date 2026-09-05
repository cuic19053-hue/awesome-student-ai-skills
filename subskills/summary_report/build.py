#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
阶段汇报 / 思想汇报 docx 生成器

格式标准：A4 纸张，页边距上下 2.54cm 左右 2.5cm；标题黑体二号居中；
称呼顶格宋体小四全角冒号；正文宋体小四 1.5 倍行距首行缩进 2 字符；
"此致"另起一行空两格，"敬礼！"另起一行顶格；落款右对齐。

使用方式：
    python build.py --data data.json --out output.docx
    python build.py --demo --out demo.docx
"""

import argparse
import json
import os
import sys
from pathlib import Path
from typing import Any, Dict, List, Optional

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

FONT_SONG = "宋体"
FONT_HEI = "黑体"
FONT_TIMES = "Times New Roman"

SIZE_ER = Pt(22)            # 二号
SIZE_XIAO_SI = Pt(12)       # 小四

def set_run_font(run, font_name: str = FONT_SONG, size=SIZE_XIAO_SI, bold: bool = False):
    run.font.name = font_name
    run.font.size = size
    run.bold = bold
    rPr = run._r.get_or_add_rPr()
    rFonts = OxmlElement('w:rFonts')
    rFonts.set(qn('w:ascii'), font_name if font_name != FONT_SONG else FONT_TIMES)
    rFonts.set(qn('w:hAnsi'), font_name if font_name != FONT_SONG else FONT_TIMES)
    rFonts.set(qn('w:eastAsia'), font_name)
    rPr.append(rFonts)

def create_document() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)
    return doc

def build_summary_report_doc(data: Dict[str, Any], output_path: str):
    doc = create_document()
    
    # 标题
    title_text = data.get("title", "阶段思想汇报")
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.space_after = Pt(12)
    run = title_p.add_run(title_text)
    set_run_font(run, FONT_HEI, SIZE_ER, bold=True)
    
    # 称呼
    salutation_p = doc.add_paragraph()
    salutation_p.paragraph_format.space_after = Pt(6)
    salutation = data.get("salutation", "尊敬的党组织：")
    run = salutation_p.add_run(salutation)
    set_run_font(run, FONT_SONG, SIZE_XIAO_SI, bold=True)
    
    # 正文段落
    paragraphs = data.get("paragraphs", [])
    for text in paragraphs:
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        p.paragraph_format.first_line_indent = Pt(24) # 2字符
        p.paragraph_format.space_after = Pt(6)
        r = p.add_run(text)
        set_run_font(r, FONT_SONG, SIZE_XIAO_SI)
        
    # 致敬语
    cz_p = doc.add_paragraph()
    cz_p.paragraph_format.first_line_indent = Pt(24)
    r = cz_p.add_run("此致")
    set_run_font(r, FONT_SONG, SIZE_XIAO_SI)
    
    jl_p = doc.add_paragraph()
    r = jl_p.add_run("敬礼！")
    set_run_font(r, FONT_SONG, SIZE_XIAO_SI, bold=True)
    
    # 落款
    signer_p = doc.add_paragraph()
    signer_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    signer_p.paragraph_format.space_before = Pt(12)
    reporter = data.get("reporter_name", "汇报人：李明")
    date_str = data.get("report_date", "2026年  月  日")
    r1 = signer_p.add_run(f"{reporter}\n{date_str}")
    set_run_font(r1, FONT_SONG, SIZE_XIAO_SI)
    
    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
    doc.save(output_path)
    print(f"成功生成阶段思想汇报: {output_path}")

def main():
    parser = argparse.ArgumentParser(description="阶段汇报 docx 生成器")
    parser.add_argument("--data", help="JSON 数据文件路径")
    parser.add_argument("--out", default="阶段思想汇报.docx", help="输出 docx 文件路径")
    parser.add_argument("--demo", action="store_true", help="使用演示数据生成")
    args = parser.parse_args()
    
    demo_data = {
        "title": "2026年第三季度思想汇报",
        "salutation": "尊敬的党组织：",
        "reporter_name": "汇报人：李明",
        "report_date": "2026年09月30日",
        "paragraphs": [
            "自被确定为入党积极分子以来，在党组织的悉心培养教育和支部同志的热情帮助下，我时刻以党员标准严格要求自己。现将近期在思想、学习、工作及作风方面的具体表现向党组织作详细汇报。",
            "在思想政治方面，我坚持读原著、学原文、悟原理，系统学习了党的最新理论成果，不断提高政治站位与政治敏锐性，深刻认识到作为新时代青年肩负的使命与担当。",
            "在专业学习方面，我保持勤奋刻苦的态度，本季度顺利完成了各项核心课程与实验研究任务，学习成绩保持在专业前 5%。同时积极参与课题组研究，将理论知识应用于实际科研攻关中。",
            "在日常工作与社会服务方面，我认真履行班长职责，积极组织开展主题团日与志愿服务活动，主动关心帮助困难同学，密切联系群众，切实发挥模范带头作用。",
            "总结近期的表现，我深知自身在理论深度的系统性上仍需进一步加强。在今后的学习工作中，我将更加严格地要求自己，虚心向优秀党员看齐，请党组织在实践中继续检验我。"
        ]
    }
    
    if args.demo or not args.data:
        build_summary_report_doc(demo_data, args.out)
    else:
        with open(args.data, "r", encoding="utf-8") as f:
            data = json.load(f)
        build_summary_report_doc(data, args.out)

if __name__ == "__main__":
    main()
