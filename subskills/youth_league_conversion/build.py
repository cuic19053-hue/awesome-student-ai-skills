#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
转正申请书 docx 生成器

格式标准：A4 纸张，页边距上下 2.54cm 左右 2.5cm；标题黑体二号居中；
称呼顶格宋体小四全角冒号；正文宋体小四 1.5 倍行距首行缩进 2 字符；
"此致"另起一行空两格，"敬礼！"另起一行顶格；落款右对齐。

使用方式：
    python build.py --data data.json --out output.docx
    python build.py --demo --out demo.docx
"""

import argparse
import json
import os
import sys
from pathlib import Path
from typing import Any, Dict, List, Optional

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

FONT_SONG = "宋体"
FONT_HEI = "黑体"
FONT_TIMES = "Times New Roman"

SIZE_ER = Pt(22)            # 二号
SIZE_XIAO_SI = Pt(12)       # 小四

def set_run_font(run, font_name: str = FONT_SONG, size=SIZE_XIAO_SI, bold: bool = False):
    run.font.name = font_name
    run.font.size = size
    run.bold = bold
    rPr = run._r.get_or_add_rPr()
    rFonts = OxmlElement('w:rFonts')
    rFonts.set(qn('w:ascii'), font_name if font_name != FONT_SONG else FONT_TIMES)
    rFonts.set(qn('w:hAnsi'), font_name if font_name != FONT_SONG else FONT_TIMES)
    rFonts.set(qn('w:eastAsia'), font_name)
    rPr.append(rFonts)

def create_document() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)
    return doc

def build_conversion_doc(data: Dict[str, Any], output_path: str):
    doc = create_document()
    
    # 标题
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.space_after = Pt(12)
    run = title_p.add_run("转正申请书")
    set_run_font(run, FONT_HEI, SIZE_ER, bold=True)
    
    # 称呼
    salutation_p = doc.add_paragraph()
    salutation_p.paragraph_format.space_after = Pt(6)
    salutation = data.get("salutation", "尊敬的党组织：")
    run = salutation_p.add_run(salutation)
    set_run_font(run, FONT_SONG, SIZE_XIAO_SI, bold=True)
    
    # 正文段落
    paragraphs = data.get("paragraphs", [])
    for text in paragraphs:
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        p.paragraph_format.first_line_indent = Pt(24)
        p.paragraph_format.space_after = Pt(6)
        r = p.add_run(text)
        set_run_font(r, FONT_SONG, SIZE_XIAO_SI)
        
    # 致敬语
    cz_p = doc.add_paragraph()
    cz_p.paragraph_format.first_line_indent = Pt(24)
    r = cz_p.add_run("此致")
    set_run_font(r, FONT_SONG, SIZE_XIAO_SI)
    
    jl_p = doc.add_paragraph()
    r = jl_p.add_run("敬礼！")
    set_run_font(r, FONT_SONG, SIZE_XIAO_SI, bold=True)
    
    # 落款
    signer_p = doc.add_paragraph()
    signer_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    signer_p.paragraph_format.space_before = Pt(12)
    applicant = data.get("applicant_name", "申请人：王伟")
    date_str = data.get("application_date", "2026年  月  日")
    r1 = signer_p.add_run(f"{applicant}\n{date_str}")
    set_run_font(r1, FONT_SONG, SIZE_XIAO_SI)
    
    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
    doc.save(output_path)
    print(f"成功生成转正申请书: {output_path}")

def main():
    parser = argparse.ArgumentParser(description="转正申请书 docx 生成器")
    parser.add_argument("--data", help="JSON 数据文件路径")
    parser.add_argument("--out", default="转正申请书.docx", help="输出 docx 文件路径")
    parser.add_argument("--demo", action="store_true", help="使用演示数据生成")
    args = parser.parse_args()
    
    demo_data = {
        "salutation": "尊敬的党组织：",
        "applicant_name": "申请人：王伟",
        "application_date": "2026年10月10日",
        "paragraphs": [
            "2025年10月10日，经党支部大会讨论表决、上级党组织批准，我光荣地成为了一名中共预备党员。截至2026年10月10日，我的预备期已满一年。今天，我郑重向党组织提出按期转正申请。",
            "在这一年的预备期内，在党组织的悉心培养与同志们的热情帮助下，我时刻以合格党员的标准严格要求自己。在思想政治上，深刻领会党的新时代路线方针政策，不断锤炼党性修养；在专业学习上，努力钻研知识，成绩保持在班级前列；在社会服务上，积极发挥先锋模范作用。",
            "经过预备期的考验，我进一步明确了入党动机，纠正了过去在工作中急于求成、不够细致的缺点。在今后的工作和学习中，我将继续保持谦虚谨慎、戒骄戒燥的作风。",
            "如果党组织批准我按期转正，我将以此为新的起点，时刻牢记入党誓词，为党和人民的事业奋斗终身；如果党组织认为我还不够成熟，决定延期转正，我也决不气馁，必将诚恳接受党组织的考验，以实际行动克服缺点，争取早日成为一名正式的共产党员。"
        ]
    }
    
    if args.demo or not args.data:
        build_conversion_doc(demo_data, args.out)
    else:
        with open(args.data, "r", encoding="utf-8") as f:
            data = json.load(f)
        build_conversion_doc(data, args.out)

if __name__ == "__main__":
    main()
